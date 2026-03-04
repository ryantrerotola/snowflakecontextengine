"""
Document Processor
==================
Handles parsing and text extraction from uploaded files:
- PDFs via Snowflake CORTEX.PARSE_DOCUMENT()
- PowerPoints via python-pptx
- Word docs via python-docx
- Plain text and Markdown via direct read
"""

import io
import tempfile
import os

from lib.snowflake_connection import get_session, call_cortex_summarize


def process_file(uploaded_file, stage_path: str) -> dict:
    """Process an uploaded file and extract text content.

    Args:
        uploaded_file: Streamlit UploadedFile object
        stage_path: Path where file was staged in Snowflake

    Returns:
        dict with keys: extracted_text, summary, status, error
    """
    file_type = _get_file_type(uploaded_file.name)

    try:
        if file_type == "pdf":
            text = _extract_pdf(stage_path)
        elif file_type == "pptx":
            text = _extract_pptx(uploaded_file)
        elif file_type == "docx":
            text = _extract_docx(uploaded_file)
        elif file_type in ("txt", "md", "csv"):
            text = _extract_text(uploaded_file)
        else:
            return {
                "extracted_text": None,
                "summary": None,
                "status": "FAILED",
                "error": f"Unsupported file type: {file_type}",
            }

        # Generate summary using Cortex
        summary = ""
        if text and len(text.strip()) > 50:
            try:
                summary = call_cortex_summarize(text[:50000])  # Limit input size
            except Exception as e:
                summary = f"(Summary generation failed: {e})"

        return {
            "extracted_text": text,
            "summary": summary,
            "status": "COMPLETED",
            "error": None,
        }

    except Exception as e:
        return {
            "extracted_text": None,
            "summary": None,
            "status": "FAILED",
            "error": str(e),
        }


def stage_file(uploaded_file, session_id: str) -> str:
    """Upload a file to the Snowflake internal stage.

    Returns the stage path.
    """
    session = get_session()
    stage_path = f"@CONTEXT_ENGINE.UPLOADS/{session_id}/{uploaded_file.name}"

    # Write to a temp file, then PUT to stage
    with tempfile.NamedTemporaryFile(delete=False, suffix=f"_{uploaded_file.name}") as tmp:
        tmp.write(uploaded_file.getvalue())
        tmp_path = tmp.name

    try:
        session.file.put(
            tmp_path,
            f"@CONTEXT_ENGINE.UPLOADS/{session_id}/",
            auto_compress=False,
            overwrite=True,
        )
    finally:
        os.unlink(tmp_path)

    return stage_path


def _extract_pdf(stage_path: str) -> str:
    """Extract text from a PDF using Snowflake CORTEX.PARSE_DOCUMENT()."""
    session = get_session()
    # PARSE_DOCUMENT returns structured content from staged files
    result = session.sql(
        f"""SELECT SNOWFLAKE.CORTEX.PARSE_DOCUMENT(
                BUILD_SCOPED_FILE_URL('{stage_path}'),
                '{{"mode": "LAYOUT"}}'
            ) AS parsed"""
    ).collect()

    if result:
        parsed = result[0]["PARSED"]
        # PARSE_DOCUMENT returns a VARIANT; extract the text content
        if isinstance(parsed, dict):
            # Extract text from all pages
            pages = parsed.get("content", [])
            if isinstance(pages, list):
                return "\n\n".join(
                    page.get("text", "") for page in pages if isinstance(page, dict)
                )
            return str(parsed)
        return str(parsed)
    return ""


def _extract_pptx(uploaded_file) -> str:
    """Extract text from a PowerPoint file."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(uploaded_file.getvalue()))
    texts = []

    for slide_num, slide in enumerate(prs.slides, 1):
        slide_texts = [f"--- Slide {slide_num} ---"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        slide_texts.append(text)
            if shape.has_table:
                table = shape.table
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip() for cell in row.cells
                    )
                    if row_text.strip(" |"):
                        slide_texts.append(row_text)
        # Include slide notes
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                slide_texts.append(f"[Speaker Notes: {notes}]")

        texts.append("\n".join(slide_texts))

    return "\n\n".join(texts)


def _extract_docx(uploaded_file) -> str:
    """Extract text from a Word document."""
    from docx import Document

    doc = Document(io.BytesIO(uploaded_file.getvalue()))
    texts = []

    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            texts.append(text)

    # Extract tables
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells)
            if row_text.strip(" |"):
                texts.append(row_text)

    return "\n".join(texts)


def _extract_text(uploaded_file) -> str:
    """Extract text from a plain text file."""
    return uploaded_file.getvalue().decode("utf-8", errors="replace")


def _get_file_type(filename: str) -> str:
    """Get normalized file type from filename."""
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return ext
